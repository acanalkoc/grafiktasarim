import csv
import math
from pathlib import Path
from typing import Optional, Any
import numpy as np

try:
    import pandas as pd
except ImportError:
    pd = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import xlrd
except ImportError:
    xlrd = None

from core.constants import SEMBOL_LISTESI
from core.models import GrafikliSeri

# -----------------------------------------------------------------------------
# Sayı ve Veri Dönüştürme Yardımcıları
# -----------------------------------------------------------------------------

def _sayiya_cevir(deger: Any) -> Optional[float]:
    """Virgüllü (3,14) ve noktalı (3.14) bilimsel sayıları güvenle ayrıştırır."""
    if deger is None or deger == "":
        return None
    if isinstance(deger, bool):
        return None
    if isinstance(deger, (int, float)):
        v = float(deger)
        return v if math.isfinite(v) else None

    metin = str(deger).strip().replace("\u00a0", "").replace(" ", "")
    if not metin:
        return None

    try:
        value = float(metin)
        return value if math.isfinite(value) else None
    except ValueError:
        pass

    if "," in metin and "." not in metin:
        try:
            value = float(metin.replace(",", "."))
            return value if math.isfinite(value) else None
        except ValueError:
            pass

    if "," in metin and "." in metin:
        if metin.rfind(",") > metin.rfind("."):
            try:
                value = float(metin.replace(".", "").replace(",", "."))
                return value if math.isfinite(value) else None
            except ValueError:
                pass
        else:
            try:
                value = float(metin.replace(",", ""))
                return value if math.isfinite(value) else None
            except ValueError:
                pass

    return None


def _istege_bagli_float(metin: str) -> Optional[float]:
    t = str(metin).strip()
    return _sayiya_cevir(t) if t else None


# -----------------------------------------------------------------------------
# Evrensel Veri Motoru (DataEngine)
# -----------------------------------------------------------------------------

class DataEngine:
    """Excel (.xlsx, .xls), CSV, TSV ve TXT dosyalarını okuma ve kaydetme motoru."""

    @staticmethod
    def get_sheet_names(path: str) -> list[str]:
        p = Path(path)
        suffix = p.suffix.lower()
        if suffix in (".xlsx", ".xlsm", ".xltx", ".xltm"):
            if openpyxl:
                wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                names = wb.sheetnames
                wb.close()
                return names
            if pd:
                return pd.ExcelFile(path).sheet_names
        elif suffix in (".xls",):
            if xlrd:
                wb = xlrd.open_workbook(path, on_demand=True)
                return wb.sheet_names()
            if pd:
                return pd.ExcelFile(path).sheet_names
        return ["Sayfa1"]

    @classmethod
    def load_file(cls, path: str, sheet_name: Optional[str] = None) -> tuple[list[str], list[list[str]], list[str], str]:
        p = Path(path)
        suffix = p.suffix.lower()
        sheet_names = ["Sayfa1"]
        active_sheet = "Sayfa1"

        if suffix in (".xlsx", ".xlsm", ".xltx", ".xltm", ".xls"):
            sheet_names = cls.get_sheet_names(path)
            if not sheet_name or sheet_name not in sheet_names:
                active_sheet = sheet_names[0]
            else:
                active_sheet = sheet_name

            raw_rows: list[list[str]] = []
            if pd:
                df = pd.read_excel(path, sheet_name=active_sheet, header=None, dtype=object)
                raw_rows = [[("" if pd.isna(cell) else str(cell).strip()) for cell in row] for row in df.values.tolist()]
            elif suffix == ".xls" and xlrd:
                wb = xlrd.open_workbook(path)
                sh = wb.sheet_by_name(active_sheet)
                for r in range(sh.nrows):
                    row = [str(sh.cell_value(r, c)).strip() for c in range(sh.ncols)]
                    raw_rows.append(row)
            elif openpyxl:
                wb = openpyxl.load_workbook(path, data_only=True)
                sh = wb[active_sheet]
                for row_cells in sh.iter_rows(values_only=True):
                    row = [("" if v is None else str(v).strip()) for v in row_cells]
                    raw_rows.append(row)
                wb.close()
            else:
                raise ImportError("Excel dosyasını açmak için 'pandas', 'openpyxl' veya 'xlrd' kütüphaneleri gereklidir.")
        else:
            raw_rows = cls._read_text_rows(path)

        raw_rows = [r for r in raw_rows if any(str(v).strip() for v in r)]
        if not raw_rows:
            raise ValueError("Dosyada veri bulunamadı veya dosya boş.")

        width = max(len(r) for r in raw_rows)
        raw_rows = [list(r) + [""] * (width - len(r)) for r in raw_rows]

        first_row = raw_rows[0]
        second_row = raw_rows[1] if len(raw_rows) > 1 else None

        first_numeric_count = sum(1 for v in first_row if _sayiya_cevir(v) is not None)
        second_numeric_count = sum(1 for v in second_row if _sayiya_cevir(v) is not None) if second_row else 0

        has_header = False
        if first_numeric_count < len(first_row) * 0.5:
            has_header = True
        elif second_row and first_numeric_count == 0 and second_numeric_count > 0:
            has_header = True

        if has_header:
            headers = [str(v).strip() or f"Sütun {i+1}" for i, v in enumerate(first_row)]
            data_rows = raw_rows[1:]
        else:
            headers = [f"Sütun_{i+1}" for i in range(width)]
            data_rows = raw_rows

        if not data_rows:
            data_rows = [first_row]
            headers = [f"Sütun_{i+1}" for i in range(width)]

        return headers, data_rows, sheet_names, active_sheet

    @staticmethod
    def _read_text_rows(path: str) -> list[list[str]]:
        encodings = ["utf-8-sig", "utf-8", "cp1254", "latin1", "cp1252"]
        content = None
        for enc in encodings:
            try:
                with open(path, "r", encoding=enc) as f:
                    content = f.read()
                if content:
                    break
            except (UnicodeDecodeError, OSError):
                continue

        if content is None:
            raise ValueError("Dosya kodlaması (encoding) okunamadı.")

        sample = "\n".join(content.splitlines()[:25])
        delimiter = ","
        try:
            sniffer = csv.Sniffer()
            dialect = sniffer.sniff(sample, delimiters=",\t; ")
            delimiter = dialect.delimiter
        except Exception:
            if "\t" in sample:
                delimiter = "\t"
            elif ";" in sample and "," not in sample:
                delimiter = ";"
            elif "," in sample:
                delimiter = ","

        lines = content.splitlines()
        reader = csv.reader(lines, delimiter=delimiter)
        rows: list[list[str]] = []
        for row in reader:
            if row and any(str(c).strip() for c in row):
                rows.append([str(c).strip() for c in row])
        return rows

    @classmethod
    def save_file(cls, path: str, headers: list[str], rows: list[list[str]]) -> None:
        p = Path(path)
        suffix = p.suffix.lower()
        if suffix in (".xlsx", ".xls"):
            if pd:
                df = pd.DataFrame(rows, columns=headers)
                # Convert numeric cells individually; majority-numeric
                # columns must never erase text/category/missing markers.
                def excel_value(value):
                    numeric = _sayiya_cevir(value)
                    return numeric if numeric is not None else value
                for index in range(len(df.columns)):
                    df.isetitem(index, df.iloc[:, index].map(excel_value))
                df.to_excel(path, index=False)
            elif openpyxl and suffix == ".xlsx":
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Veri"
                ws.append(headers)
                for r in rows:
                    ws.append([(_sayiya_cevir(v) if _sayiya_cevir(v) is not None else v) for v in r])
                wb.save(path)
            else:
                raise ImportError("Excel kaydetmek için pandas veya openpyxl gereklidir.")
        else:
            with open(path, "w", encoding="utf-8-sig", newline="") as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                writer.writerows(rows)


# -----------------------------------------------------------------------------
# Sütun Rolleri ve Seri İnşa Motoru
# -----------------------------------------------------------------------------

def otomatik_rol_tahmin_et(headers: list[str], mapping_mode: str = "Auto") -> list[str]:
    """Sütun başlıklarını analiz ederek bilimsel standartlarda X, Y, yErr, xErr rollerini otomatik belirler."""
    num_cols = len(headers)
    if num_cols == 0:
        return []

    if mapping_mode == "Pairs":
        return ["X" if i % 2 == 0 else "Y" for i in range(num_cols)]
    elif mapping_mode == "Shared X":
        return ["X"] + ["Y"] * (num_cols - 1)

    roles = []
    lower_headers = [h.strip().lower() for h in headers]

    for i, h in enumerate(lower_headers):
        # Y Hata kontrolü
        import re
        tokens = set(re.findall(r"[^\W_]+", h))
        is_y_err = bool(tokens.intersection({"yerr", "hata", "error", "sd", "std", "se", "sem", "sigma", "belirsizlik"})) or "±" in h
        is_x_err = any(kw in h for kw in ["xerr", "x_err", "x error", "x_hata", "x hata"])

        if is_y_err and not is_x_err:
            roles.append("yErr")
        elif is_x_err:
            roles.append("xErr")
        elif i == 0 or any(x_kw in h for x_kw in ["wavelength", "dalga", "zaman", "time", "derişim", "conc", "gerilme", "strain", "x1", "x2", "x3", "x4"]) or h in ("x", "t"):
            roles.append("X")
        else:
            roles.append("Y")

    if not any(r == "X" for r in roles):
        roles[0] = "X"

    return roles


def _categorical_x(metadata, index):
    return metadata is not None and index is not None and 0 <= index < len(metadata.columns) and metadata.columns[index].data_type == "Categorical"


def seri_listesi_olustur(
    headers: list[str],
    rows: list[list[str]],
    roles: Optional[list[str]] = None,
    mapping_mode: str = "Auto",
    mode: Optional[str] = None,
    *,
    worksheet_id: Optional[str] = None,
    metadata: Optional[Any] = None,
) -> list[GrafikliSeri]:
    """Build plot-ready series from a worksheet's rows.

    ``worksheet_id``/``metadata`` are optional (v6.2): when supplied, each
    built series is stamped with the worksheet's stable id plus its X/Y
    column UUIDs (from ``metadata``, a ``core.worksheet.WorksheetMetadata``)
    so the binding survives renames and never collides with a same-named
    column in a different worksheet. Callers that don't pass them (existing
    single-worksheet call sites) get the pre-v6.2 behaviour unchanged.
    """
    if mode is not None:
        mapping_mode = mode
    if not headers or not rows:
        return []

    num_cols = len(headers)
    if mapping_mode in ("Pairs", "Shared X") or roles is None or len(roles) != num_cols:
        roles = otomatik_rol_tahmin_et(headers, mapping_mode=mapping_mode)

    series_list: list[GrafikliSeri] = []
    current_x_col = None
    has_x_role = any(r == "X" for r in roles)

    for c_idx, role in enumerate(roles):
        if role == "X":
            current_x_col = c_idx
        elif role == "Y":
            x_col = current_x_col if (has_x_role and current_x_col is not None) else None
            y_col = c_idx
            yerr_col = None
            xerr_col = None

            if c_idx + 1 < num_cols and roles[c_idx + 1] == "yErr":
                yerr_col = c_idx + 1
            if c_idx + 1 < num_cols and roles[c_idx + 1] == "xErr":
                xerr_col = c_idx + 1
            elif c_idx + 2 < num_cols and roles[c_idx + 2] == "xErr":
                xerr_col = c_idx + 2

            xs, ys, yerrs, xerrs, row_idxs = [], [], [], [], []
            categorical = _categorical_x(metadata, x_col)
            for r_idx, row in enumerate(rows):
                if x_col is not None:
                    xv = float(r_idx + 1) if categorical else _sayiya_cevir(row[x_col] if x_col < len(row) else None)
                    if xv is None:
                        continue  # Explicit X is missing: never invent an abscissa.
                else:
                    xv = float(r_idx + 1)

                yv = _sayiya_cevir(row[y_col] if y_col < len(row) else None)
                if yv is not None:
                    xs.append(xv)
                    ys.append(yv)
                    row_idxs.append(r_idx)
                    if yerr_col is not None:
                        ev = _sayiya_cevir(row[yerr_col] if yerr_col < len(row) else None)
                        yerrs.append(ev if ev is not None and ev >= 0 else np.nan)
                    if xerr_col is not None:
                        xev = _sayiya_cevir(row[xerr_col] if xerr_col < len(row) else None)
                        xerrs.append(xev if xev is not None and xev >= 0 else np.nan)

            if xs and ys:
                s_name = headers[y_col] if headers[y_col] else f"Seri {len(series_list) + 1}"
                x_arr = np.asarray(xs, dtype=float)
                y_arr = np.asarray(ys, dtype=float)
                yerr_arr = np.asarray(yerrs, dtype=float) if (yerr_col is not None and yerrs) else None
                xerr_arr = np.asarray(xerrs, dtype=float) if (xerr_col is not None and xerrs) else None
                x_hdr = headers[x_col] if x_col is not None else "İndeks"

                s = GrafikliSeri(
                    name=s_name,
                    x=x_arr,
                    y=y_arr,
                    yerr=yerr_arr,
                    xerr=xerr_arr,
                    x_header=x_hdr,
                    x_labels=[str(rows[i][x_col]) if x_col < len(rows[i]) else "" for i in row_idxs] if categorical else None,
                    y_header=headers[y_col],
                    marker=SEMBOL_LISTESI[len(series_list) % len(SEMBOL_LISTESI)],
                    x_col_idx=x_col,
                    y_col_idx=y_col,
                    yerr_col_idx=yerr_col,
                    yerr_column_id=metadata.legacy_column_id(yerr_col) if metadata is not None and yerr_col is not None else None,
                    xerr_col_idx=xerr_col,
                    xerr_column_id=metadata.legacy_column_id(xerr_col) if metadata is not None and xerr_col is not None else None,
                    # Stable, column-based identity (not the display name, so
                    # two same-named columns never collide) + the original
                    # worksheet row for every kept (x, y) pair, so callers
                    # (Analysis Studio) can scatter row-aligned results back
                    # to the correct rows even though blank-Y rows were
                    # skipped/compacted above.
                    series_uid=metadata.legacy_column_id(y_col) if metadata is not None else f"col{y_col}",
                    row_indices=np.asarray(row_idxs, dtype=int),
                    worksheet_id=worksheet_id,
                    x_column_id=metadata.legacy_column_id(x_col) if (metadata is not None and x_col is not None) else None,
                    y_column_id=metadata.legacy_column_id(y_col) if metadata is not None else None,
                )
                series_list.append(s)

    return series_list


def series_from_columns(
    headers: list[str],
    rows: list[list[str]],
    x_idx: Optional[int],
    y_idx: int,
    *,
    name: Optional[str] = None,
    worksheet_id: Optional[str] = None,
    metadata: Optional[Any] = None,
) -> Optional[GrafikliSeri]:
    """Build a single XY series from explicit column indices.

    Used by the v6.2 layer data-assignment flow (worksheet -> X column -> Y
    column -> target layer, see V62 requirements §3), where the user picks
    columns directly rather than relying on role auto-detection. Stamps
    ``worksheet_id``/column UUIDs exactly like ``seri_listesi_olustur`` so
    the resulting binding is identity-based, not positional.
    """
    if y_idx is None or y_idx >= len(headers):
        return None
    xs, ys, row_idxs = [], [], []
    categorical = _categorical_x(metadata, x_idx)
    for r_idx, row in enumerate(rows):
        if x_idx is not None:
            xv = float(r_idx + 1) if categorical else _sayiya_cevir(row[x_idx] if x_idx < len(row) else None)
            if xv is None:
                continue  # Row numbers are allowed only when no X column was chosen.
        else:
            xv = float(r_idx + 1)
        yv = _sayiya_cevir(row[y_idx] if y_idx < len(row) else None)
        if yv is not None:
            xs.append(xv)
            ys.append(yv)
            row_idxs.append(r_idx)
    if not xs:
        return None
    x_hdr = headers[x_idx] if x_idx is not None else "İndeks"
    return GrafikliSeri(
        name=name or headers[y_idx] or f"Seri {y_idx + 1}",
        x=np.asarray(xs, dtype=float),
        y=np.asarray(ys, dtype=float),
        x_header=x_hdr,
        x_labels=[str(rows[i][x_idx]) if x_idx < len(rows[i]) else "" for i in row_idxs] if categorical else None,
        y_header=headers[y_idx],
        marker=SEMBOL_LISTESI[y_idx % len(SEMBOL_LISTESI)],
        x_col_idx=x_idx,
        y_col_idx=y_idx,
        series_uid=metadata.legacy_column_id(y_idx) if metadata is not None else f"col{y_idx}",
        row_indices=np.asarray(row_idxs, dtype=int),
        worksheet_id=worksheet_id,
        x_column_id=metadata.legacy_column_id(x_idx) if (metadata is not None and x_idx is not None) else None,
        y_column_id=metadata.legacy_column_id(y_idx) if metadata is not None else None,
    )


build_xy_series = seri_listesi_olustur
XYSeries = GrafikliSeri
